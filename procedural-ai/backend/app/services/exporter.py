"""Экспорт документа в DOCX и PDF с сохранением текстовой редактируемости.

Исходное представление документа — ограниченное подмножество HTML,
формируемое генератором и редактором (p, h1–h3, b/i/u, br, ul/ol/li,
table, div.addressee, div.consent, span)."""
from __future__ import annotations

import io

import re
from html import unescape
from html.parser import HTMLParser
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from ..config import settings

BLOCK_TAGS = {"p", "h1", "h2", "h3", "li"}


# --- разбор HTML в простую модель ------------------------------------------------

class Run:
    __slots__ = ("text", "bold", "italic", "underline", "gap", "br")

    def __init__(self, text="", bold=False, italic=False, underline=False, gap=False, br=False):
        self.text, self.bold, self.italic, self.underline, self.gap, self.br = text, bold, italic, underline, gap, br


class Block:
    def __init__(self, kind: str, classes: set[str], container: str | None):
        self.kind = kind  # p | h1 | h2 | h3 | li
        self.classes = classes
        self.container = container  # addressee | consent | None
        self.runs: list[Run] = []
        self.list_kind: str | None = None
        self.list_index = 0


class Table:
    def __init__(self, classes: set[str]):
        self.classes = classes
        self.rows: list[list[list[Block]]] = []


class _Parser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.items: list[Block | Table] = []
        self.fmt = []  # стек (tag, attrs)
        self.cur: Block | None = None
        self.container_stack: list[str | None] = []
        self.table: Table | None = None
        self.cell: list[Block] | None = None
        self.list_stack: list[list] = []

    def _classes(self, attrs) -> set[str]:
        return set((dict(attrs).get("class") or "").split())

    def _container(self):
        for c in reversed(self.container_stack):
            if c:
                return c
        return None

    def _target(self) -> list:
        return self.cell if self.cell is not None else self.items

    def _ensure_block(self):
        if self.cur is None:
            self.cur = Block("p", set(), self._container())
            self._target().append(self.cur)

    def handle_starttag(self, tag, attrs):
        cls = self._classes(attrs)
        if tag == "div":
            self.container_stack.append("addressee" if "addressee" in cls else "consent" if "consent" in cls else None)
        elif tag in ("p", "h1", "h2", "h3"):
            self.cur = Block(tag, cls, self._container())
            self._target().append(self.cur)
        elif tag in ("ul", "ol"):
            self.list_stack.append([tag, 0])
        elif tag == "li":
            self.cur = Block("li", cls, self._container())
            if self.list_stack:
                self.list_stack[-1][1] += 1
                self.cur.list_kind, self.cur.list_index = self.list_stack[-1][0], self.list_stack[-1][1]
            self._target().append(self.cur)
        elif tag == "table":
            self.table = Table(cls)
            self.items.append(self.table)
        elif tag == "tr" and self.table is not None:
            self.table.rows.append([])
        elif tag in ("td", "th") and self.table is not None:
            self.cell = []
            self.table.rows[-1].append(self.cell)
            self.cur = Block("p", self._classes(attrs) | ({"bold"} if tag == "th" else set()), None)
            self.cell.append(self.cur)
        elif tag == "br":
            self._ensure_block()
            self.cur.runs.append(Run(br=True))
        elif tag in ("b", "strong", "i", "em", "u", "span"):
            self.fmt.append((tag, cls))

    def handle_endtag(self, tag):
        if tag == "div":
            if self.container_stack:
                self.container_stack.pop()
        elif tag in ("p", "h1", "h2", "h3", "li"):
            self.cur = None
        elif tag in ("ul", "ol"):
            if self.list_stack:
                self.list_stack.pop()
        elif tag in ("td", "th"):
            self.cell = None
            self.cur = None
        elif tag == "table":
            self.table = None
        elif tag in ("b", "strong", "i", "em", "u", "span"):
            for i in range(len(self.fmt) - 1, -1, -1):
                if self.fmt[i][0] == tag:
                    del self.fmt[i]
                    break

    def handle_data(self, data):
        if not data:
            return
        if self.cur is None:
            if not data.strip():
                return
            self._ensure_block()
        tags = {t for t, _ in self.fmt}
        gap = any("gap" in c for _, c in self.fmt)
        pre = "pre" in self.cur.classes
        text = data if pre else re.sub(r"\s+", " ", data)
        self.cur.runs.append(Run(text, bold=bool(tags & {"b", "strong"}) or "bold" in self.cur.classes,
                                 italic=bool(tags & {"i", "em"}), underline="u" in tags, gap=gap))


def parse(html: str) -> list[Block | Table]:
    p = _Parser()
    p.feed(html)
    p.close()
    return p.items


def _block_text(b: Block) -> str:
    return "".join("\n" if r.br else r.text for r in b.runs).strip()


# --- DOCX ---------------------------------------------------------------------

def _add_page_number(section):
    header = section.header
    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(12)


def _no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def _fill_paragraph(par, block: Block, base_size=14):
    for r in block.runs:
        if r.br:
            par.add_run().add_break()
            continue
        run = par.add_run(r.text)
        run.bold = r.bold or block.kind in ("h1", "h2") or None
        run.italic = r.italic or None
        run.underline = r.underline or None
        run.font.size = Pt(12 if "small" in block.classes else base_size)
        if r.gap:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _format_paragraph(par, block: Block):
    pf = par.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    cls = block.classes
    if block.kind in ("h1", "h2", "h3") or "center" in cls or "doc-subtitle" in cls:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if block.kind == "h1":
            pf.space_before = Pt(6)
    elif "right" in cls:
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif block.container == "addressee":
        pf.left_indent = Cm(8.5)
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif block.container == "consent":
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif "pre" in cls:
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(1.25)
    else:
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(1.25)


def to_docx(html: str, *, title: str, draft_mark: str | None = None) -> bytes:
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Times New Roman")
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    _add_page_number(sec)
    doc.core_properties.title = title
    doc.core_properties.comments = "Сформировано PROCEDURAL AI. Документ требует проверки уполномоченным лицом."
    if draft_mark:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(draft_mark)
        r.italic = True
        r.font.size = Pt(10)

    for item in parse(html):
        if isinstance(item, Table):
            ncols = max((len(r) for r in item.rows), default=0)
            if not ncols:
                continue
            table = doc.add_table(rows=len(item.rows), cols=ncols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            if "card" in item.classes:
                table.style = "Table Grid"
            else:
                _no_borders(table)
            for ri, row in enumerate(item.rows):
                for ci, cell_blocks in enumerate(row):
                    cell = table.cell(ri, ci)
                    first = True
                    for b in cell_blocks:
                        par = cell.paragraphs[0] if first else cell.add_paragraph()
                        first = False
                        if "right" in b.classes or "sign-name" in b.classes:
                            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        _fill_paragraph(par, b, base_size=13 if "card" in item.classes else 14)
            continue
        par = doc.add_paragraph()
        if item.kind == "li" and item.list_kind:
            prefix = f"{item.list_index}. " if item.list_kind == "ol" else "– "
            par.add_run(prefix)
        _format_paragraph(par, item)
        _fill_paragraph(par, item)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- PDF ----------------------------------------------------------------------

def _font_dir() -> tuple[Path, str, str]:
    candidates = [
        (settings.pdf_font_path, None, None),
        ("/usr/share/fonts/truetype/liberation", "LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf"),
        ("/usr/share/fonts/truetype/dejavu", "DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf"),
    ]
    for d, reg, bold in candidates:
        if not d:
            continue
        p = Path(d)
        if p.is_dir():
            reg = reg or next((f.name for f in p.glob("*Serif*Regular*.ttf")), None) or next((f.name for f in p.glob("*.ttf")), None)
            bold = bold or next((f.name for f in p.glob("*Serif*Bold.ttf")), reg)
            if reg and (p / reg).exists():
                return p, reg, bold if (p / bold).exists() else reg
    raise RuntimeError("Не найден шрифт с поддержкой кириллицы для PDF (задайте PAI_PDF_FONT)")


PDF_CSS = """
@font-face {{ font-family: docserif; src: url({reg}); }}
@font-face {{ font-family: docserif; src: url({bold}); font-weight: bold; }}
* {{ font-family: docserif; font-size: 13pt; line-height: 1.25; }}
p {{ margin: 0; text-align: justify; text-indent: 1.25cm; }}
p.center, p.doc-subtitle {{ text-align: center; text-indent: 0; }}
p.right {{ text-align: right; text-indent: 0; }}
p.small {{ font-size: 11pt; }}
h1, h2 {{ text-align: center; font-weight: bold; margin: 8pt 0 0 0; font-size: 14pt; }}
div.addressee p {{ margin-left: 8cm; text-indent: 0; text-align: left; }}
div.consent p {{ text-align: right; text-indent: 0; }}
table {{ width: 100%; border-collapse: collapse; margin: 4pt 0; }}
table.card td {{ border: 0.5pt solid black; padding: 2pt 4pt; font-size: 12pt; }}
td.right, td.sign-name {{ text-align: right; }}
span.gap {{ background-color: #fff2a8; }}
p.draft {{ text-align: right; font-size: 9pt; font-style: italic; text-indent: 0; }}
"""


def to_pdf(html: str, *, title: str, draft_mark: str | None = None) -> bytes:
    import pymupdf as fitz

    font_dir, reg, bold = _font_dir()
    css = PDF_CSS.format(reg=reg, bold=bold)
    body = (f'<p class="draft">{draft_mark}</p>' if draft_mark else "") + html
    # Абзацы класса pre сохраняют переносы строк
    body = re.sub(r'(<p class="pre">)(.*?)(</p>)', lambda m: m.group(1) + m.group(2).replace("\n", "<br>") + m.group(3),
                  body, flags=re.S)
    story = fitz.Story(html=f"<html><body>{body}</body></html>", user_css=css, archive=fitz.Archive(str(font_dir)))
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    mediabox = fitz.paper_rect("a4")
    # Поля: левое 30 мм, правое 15 мм, верхнее и нижнее 20 мм
    where = fitz.Rect(85, 57, mediabox.width - 42, mediabox.height - 57)
    more = 1
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    pdf = fitz.open("pdf", buf.getvalue())
    font_file = str(font_dir / reg)
    if pdf.page_count > 1:
        for i, page in enumerate(pdf, start=1):
            page.insert_text((mediabox.width / 2 - 5, 35), str(i), fontsize=11, fontname="docserif", fontfile=font_file)
    pdf.set_metadata({"title": title, "creator": "PROCEDURAL AI", "producer": "PROCEDURAL AI",
                      "subject": "Проект процессуального документа"})
    try:
        pdf.subset_fonts()
    except Exception:  # noqa: BLE001 - уменьшение размера необязательно
        pass
    out = pdf.tobytes(deflate=True, garbage=3)
    pdf.close()
    return out


def html_to_plain(html: str) -> str:
    lines = []
    for item in parse(html):
        if isinstance(item, Table):
            for row in item.rows:
                lines.append(" | ".join(" ".join(_block_text(b) for b in cell) for cell in row))
        else:
            lines.append(_block_text(item))
    return unescape("\n".join(lines))


__all__ = ["to_docx", "to_pdf", "html_to_plain", "parse"]

