"""Извлечение текста из материалов и разбиение на логические фрагменты."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from . import asr, ocr
from .file_safety import IMAGE_EXTS, MEDIA_EXTS, decode_text

MAX_FRAGMENT = 1500


@dataclass
class Fragment:
    text: str
    page: int | None = None
    paragraph: int | None = None
    method: str = "digital"
    confidence: float | None = None


@dataclass
class ExtractionResult:
    fragments: list[Fragment] = field(default_factory=list)
    page_count: int | None = None
    method: str = "digital"
    machine_generated: bool = False
    warnings: list[str] = field(default_factory=list)
    needs_ocr: bool = False


def _split_long(text: str) -> list[str]:
    text = text.strip()
    if len(text) <= MAX_FRAGMENT:
        return [text] if text else []
    parts, buf = [], ""
    for sent in re.split(r"(?<=[.!?;])\s+", text):
        if len(buf) + len(sent) + 1 > MAX_FRAGMENT and buf:
            parts.append(buf.strip())
            buf = ""
        while len(sent) > MAX_FRAGMENT:
            parts.append(sent[:MAX_FRAGMENT])
            sent = sent[MAX_FRAGMENT:]
        buf += sent + " "
    if buf.strip():
        parts.append(buf.strip())
    return parts


def _paragraphs(text: str, page: int | None = None, method: str = "digital") -> list[Fragment]:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks = re.split(r"\n\s*\n", text)
    if len(blocks) == 1:
        blocks = text.split("\n")
    out, n = [], 0
    for block in blocks:
        clean = re.sub(r"[ \t]+", " ", block.replace("\n", " ")).strip()
        for piece in _split_long(clean):
            n += 1
            out.append(Fragment(piece, page=page, paragraph=n, method=method))
    return out


def _docx(data: bytes) -> ExtractionResult:
    import docx

    doc = docx.Document(io.BytesIO(data))
    res = ExtractionResult()
    n = 0
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = "".join(t.text or "" for t in child.iter() if t.tag.endswith("}t")).strip()
            for piece in _split_long(text):
                n += 1
                res.fragments.append(Fragment(piece, paragraph=n))
        elif tag == "tbl":
            for row in child.iter():
                if not row.tag.endswith("}tr"):
                    continue
                cells = []
                for cell in row.iterchildren():
                    if cell.tag.endswith("}tc"):
                        cells.append(" ".join((t.text or "") for t in cell.iter() if t.tag.endswith("}t")).strip())
                line = " | ".join(c for c in cells if c)
                if line:
                    n += 1
                    res.fragments.append(Fragment(line, paragraph=n))
    return res


def _pdf(data: bytes) -> ExtractionResult:
    import pymupdf as fitz

    res = ExtractionResult(method="digital")
    with fitz.open(stream=data, filetype="pdf") as pdf:
        res.page_count = pdf.page_count
        ocr_pages = 0
        for pno, page in enumerate(pdf, start=1):
            blocks = [b for b in page.get_text("blocks") if b[6] == 0 and b[4].strip()]
            if blocks:
                para = 0
                for b in blocks:
                    for piece in _split_long(re.sub(r"\s+", " ", b[4])):
                        para += 1
                        res.fragments.append(Fragment(piece, page=pno, paragraph=para))
                continue
            # Страница без текстового слоя — требуется OCR.
            png = page.get_pixmap(dpi=300).tobytes("png")
            try:
                _engine, lines = ocr.recognize(png)
            except ocr.OcrUnavailable:
                res.needs_ocr = True
                continue
            ocr_pages += 1
            for i, line in enumerate([ln for ln in lines if ln.text.strip()], start=1):
                res.fragments.append(Fragment(line.text, page=pno, paragraph=i, method="ocr", confidence=line.confidence))
        if ocr_pages:
            res.machine_generated = True
            res.method = "mixed" if ocr_pages < res.page_count else "ocr"
        if res.needs_ocr:
            res.warnings.append("Часть страниц не содержит текстового слоя, локальный OCR недоступен")
    return res


def _image(data: bytes) -> ExtractionResult:
    res = ExtractionResult(method="ocr", machine_generated=True, page_count=1)
    try:
        engine, lines = ocr.recognize(data)
    except ocr.OcrUnavailable as exc:
        res.needs_ocr = True
        res.method = "none"
        res.warnings.append(str(exc))
        return res
    res.warnings.append(f"Текст распознан машинным способом ({engine}); требуется сверка с оригиналом")
    for i, line in enumerate([ln for ln in lines if ln.text.strip()], start=1):
        res.fragments.append(Fragment(line.text, page=1, paragraph=i, method="ocr", confidence=line.confidence))
    return res


_RTF_SKIP = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "object", "header", "footer",
             "listtable", "listoverridetable", "generator", "themedata", "colorschememapping",
             "datastore", "latentstyles", "rsidtbl", "xmlnstbl", "mmathPr", "pgdsctbl"}


def rtf_to_text(raw: str) -> str:
    """Упрощённый разбор RTF: извлекается текст без форматирования."""
    out: list[str] = []
    stack: list[bool] = []
    skip = False
    i, n = 0, len(raw)
    uc_skip = 1
    pending_bytes = bytearray()

    def flush():
        if pending_bytes:
            out.append(pending_bytes.decode("cp1251", errors="replace"))
            pending_bytes.clear()

    while i < n:
        ch = raw[i]
        if ch == "{":
            stack.append(skip)
            i += 1
        elif ch == "}":
            flush()
            skip = stack.pop() if stack else False
            i += 1
        elif ch == "\\":
            i += 1
            if i >= n:
                break
            c = raw[i]
            if c in "\\{}":
                flush()
                if not skip:
                    out.append(c)
                i += 1
            elif c == "'":
                hexv = raw[i + 1 : i + 3]
                i += 3
                if not skip:
                    try:
                        pending_bytes.append(int(hexv, 16))
                    except ValueError:
                        pass
            elif c == "*":
                skip = True
                i += 1
            elif c.isalpha():
                m = re.match(r"([a-zA-Z]+)(-?\d+)? ?", raw[i:])
                word, arg = m.group(1), m.group(2)
                i += m.end()
                flush()
                if word in _RTF_SKIP:
                    skip = True
                elif word in ("par", "line", "row"):
                    if not skip:
                        out.append("\n")
                elif word == "tab" or word == "cell":
                    if not skip:
                        out.append("\t")
                elif word == "uc" and arg:
                    uc_skip = int(arg)
                elif word == "u" and arg:
                    code = int(arg)
                    if code < 0:
                        code += 65536
                    if not skip:
                        out.append(chr(code))
                    # пропуск символов-заменителей
                    j = 0
                    while j < uc_skip and i < n:
                        if raw[i] == "\\" and i + 1 < n and raw[i + 1] == "'":
                            i += 4
                        else:
                            i += 1
                        j += 1
            else:
                i += 1
        else:
            flush()
            if not skip and ch not in "\r\n":
                out.append(ch)
            i += 1
    flush()
    return "".join(out)


def _xlsx(data: bytes) -> ExtractionResult:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    res = ExtractionResult()
    n = 0
    max_rows = 20000
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                n += 1
                res.fragments.append(Fragment(" | ".join(cells), paragraph=n))
            if n >= max_rows:
                res.warnings.append(f"Импорт ограничен {max_rows} строками")
                return res
    wb.close()
    return res


def _csv(data: bytes) -> ExtractionResult:
    text = decode_text(data)
    res = ExtractionResult()
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect), start=1):
        line = " | ".join(c.strip() for c in row if c.strip())
        if line:
            res.fragments.append(Fragment(line, paragraph=i))
        if i >= 20000:
            res.warnings.append("Импорт ограничен 20000 строками")
            break
    return res


def _media(data: bytes, ext: str) -> ExtractionResult:
    res = ExtractionResult(method="asr", machine_generated=True)
    try:
        segments = asr.transcribe(data, ext)
    except asr.AsrUnavailable as exc:
        res.method = "none"
        res.warnings.append(str(exc) + "; материал сохранён без расшифровки")
        return res
    for i, (text, conf, start) in enumerate(segments, start=1):
        mm, ss = divmod(int(start), 60)
        res.fragments.append(Fragment(f"[{mm:02d}:{ss:02d}] {text}", paragraph=i, method="asr", confidence=conf))
    res.warnings.append("Расшифровка сформирована машинным способом и требует сверки с записью")
    return res


def extract(data: bytes, ext: str) -> ExtractionResult:
    if ext == "txt":
        res = ExtractionResult()
        res.fragments = _paragraphs(decode_text(data))
        return res
    if ext == "csv":
        return _csv(data)
    if ext == "docx":
        return _docx(data)
    if ext == "pdf":
        return _pdf(data)
    if ext == "rtf":
        res = ExtractionResult()
        res.fragments = _paragraphs(rtf_to_text(data.decode("latin-1")))
        return res
    if ext == "xlsx":
        return _xlsx(data)
    if ext in IMAGE_EXTS:
        return _image(data)
    if ext in MEDIA_EXTS:
        return _media(data, ext)
    raise ValueError(f"Нет обработчика для .{ext}")
