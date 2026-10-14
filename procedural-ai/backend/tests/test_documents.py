"""Сценарий А. Создание дела и формирование документа по шаблону."""
import io


def _victim(client, users, cid):
    return client.post(f"/api/v1/cases/{cid}/participants", headers=users["inv1"], json={
        "role": "victim", "last_name": "Петров", "first_name": "Иван", "patronymic": "Сергеевич",
        "data": {"birth_date": "1980-02-01", "address_actual": "г. Тестовск, ул. Лесная, д. 7, кв. 12"}}).json()


def test_template_generation_with_gaps(client, users, case):
    cid = case["id"]
    victim = _victim(client, users, cid)
    r = client.post(f"/api/v1/cases/{cid}/documents", headers=users["inv1"], json={"type_code": "victim_recognition"})
    assert r.status_code == 201, r.text
    doc = r.json()
    v = doc["current_version"]
    # Данные карточки дела подставлены с указанием источника
    assert v["field_values"]["case_number"] == "12601690001000045"
    assert v["field_sources"]["case_number"]["type"] == "case"
    assert v["field_values"]["victim"] == victim["id"]
    assert "12601690001000045" in v["content_html"]
    assert "Петров Иван Сергеевич" in v["content_html"]
    # Отсутствующие сведения не домысливаются
    assert 'class="gap"' in v["content_html"]
    assert doc["status"] == "awaiting_input"
    assert v["validation"]["summary"]["errors"] > 0


def test_fill_review_approve_export(client, users, case):
    cid = case["id"]
    _victim(client, users, cid)
    doc = client.post(f"/api/v1/cases/{cid}/documents", headers=users["inv1"], json={"type_code": "victim_recognition"}).json()
    values = dict(doc["current_version"]["field_values"])
    sources = dict(doc["current_version"]["field_sources"])
    values.update({"circumstances": "10.03.2026 неустановленное лицо тайно похитило имущество потерпевшего.",
                   "harm": "имущественный вред", "harm_details": "15 000 рублей", "doc_date": "2026-03-11"})
    r = client.post(f"/api/v1/documents/{doc['id']}/versions", headers=users["inv1"],
                    json={"values": values, "sources": sources})
    assert r.status_code == 201, r.text
    d = r.json()
    assert d["current_version"]["number"] == 2
    assert d["status"] == "needs_review"
    assert d["current_version"]["validation"]["summary"]["errors"] == 0
    # Отметка о проверке требует явного подтверждения
    assert client.post(f"/api/v1/documents/{doc['id']}/status", headers=users["inv1"], json={"status": "reviewed"}).status_code == 422
    r = client.post(f"/api/v1/documents/{doc['id']}/status", headers=users["inv1"], json={"status": "reviewed", "confirm": True})
    assert r.json()["status"] == "reviewed"
    r = client.post(f"/api/v1/documents/{doc['id']}/status", headers=users["inv1"], json={"status": "approved", "confirm": True})
    assert r.json()["status"] == "approved"
    r = client.post(f"/api/v1/documents/{doc['id']}/export?fmt=docx", headers=users["inv1"])
    assert r.status_code == 200
    import docx

    text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    assert "ПОСТАНОВЛЕНИЕ" in text and "ПРОЕКТ ДОКУМЕНТА" not in text
    r = client.post(f"/api/v1/documents/{doc['id']}/export?fmt=pdf", headers=users["inv1"])
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    import pymupdf

    pdf_text = pymupdf.open("pdf", r.content)[0].get_text()
    assert "ПОСТАНОВЛЕНИЕ" in pdf_text
    d = client.get(f"/api/v1/documents/{doc['id']}", headers=users["inv1"]).json()
    assert d["status"] == "exported"
    exports = client.get(f"/api/v1/documents/{doc['id']}/exports", headers=users["inv1"]).json()
    assert len(exports) == 2
    # Редактирование после утверждения сбрасывает отметку
    values["harm_details"] = "16 000 рублей"
    r = client.post(f"/api/v1/documents/{doc['id']}/versions", headers=users["inv1"], json={"values": values, "sources": sources})
    assert r.json()["status"] == "needs_review" and r.json()["approved_by"] is None
    diff = client.get(f"/api/v1/documents/{doc['id']}/diff?a=2&b=3", headers=users["inv1"]).json()
    assert any(x["op"] == "added" and "16 000" in x["text"] for x in diff["lines"])
    assert diff["fields"][0]["key"] == "harm_details"


def test_draft_export_is_marked(client, users, case):
    doc = client.post(f"/api/v1/cases/{case['id']}/documents", headers=users["inv1"], json={"type_code": "acceptance"}).json()
    r = client.post(f"/api/v1/documents/{doc['id']}/export?fmt=docx", headers=users["inv1"])
    import docx

    text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    assert "ПРОЕКТ ДОКУМЕНТА" in text


def test_versions_immutable_after_case_change(client, users, case):
    cid = case["id"]
    doc = client.post(f"/api/v1/cases/{cid}/documents", headers=users["inv1"], json={"type_code": "acceptance"}).json()
    html_before = doc["current_version"]["content_html"]
    client.patch(f"/api/v1/cases/{cid}", headers=users["inv1"], json={"case_number": "99999999999999999"})
    v1 = client.get(f"/api/v1/documents/{doc['id']}/versions/1", headers=users["inv1"]).json()
    assert v1["content_html"] == html_before
    assert "12601690001000045" in v1["content_html"]
    val = client.post(f"/api/v1/documents/{doc['id']}/validate", headers=users["inv1"]).json()
    assert any(i["code"] == "case_number" for i in val["items"])


def test_cannot_review_with_errors(client, users, case):
    doc = client.post(f"/api/v1/cases/{case['id']}/documents", headers=users["inv1"], json={"type_code": "initiation"}).json()
    client.post(f"/api/v1/documents/{doc['id']}/status", headers=users["inv1"], json={"status": "needs_review"})
    r = client.post(f"/api/v1/documents/{doc['id']}/status", headers=users["inv1"], json={"status": "reviewed", "confirm": True})
    assert r.status_code in (409,)


def test_edited_html_is_sanitized(client, users, case):
    doc = client.post(f"/api/v1/cases/{case['id']}/documents", headers=users["inv1"],
                      json={"mode": "manual", "title": "Справка (ручной документ)"}).json()
    assert doc["generation_mode"] == "manual" and doc["template"] is None
    r = client.post(f"/api/v1/documents/{doc['id']}/versions", headers=users["inv1"], json={
        "edited": True, "content_html": '<p onclick="x()">Текст<script>alert(1)</script><img src=x onerror=y></p>'})
    html = r.json()["current_version"]["content_html"]
    assert "script" not in html and "onclick" not in html and "img" not in html and "Текст" in html


def test_kit_generation(client, users, case):
    r = client.post(f"/api/v1/cases/{case['id']}/kits", headers=users["inv1"], json={"kit_code": "detention_kit"})
    assert r.status_code == 201
    docs = r.json()["documents"]
    assert len(docs) == 4 and len({d["id"] for d in docs}) == 4
    assert all(d["kit_id"] == r.json()["kit_id"] for d in docs)


def test_assistant_suggest(client, users, case):
    r = client.post(f"/api/v1/cases/{case['id']}/assistant/suggest", headers=users["inv1"],
                    json={"text": "нужно назначить судебную экспертизу по изъятому телефону"})
    s = r.json()["suggestions"]
    assert s[0]["code"] == "expertise_order"
    assert any(q["key"] == "questions" for q in s[0]["clarifying_questions"])


def test_select_value_validation(client, users, case):
    r = client.post(f"/api/v1/cases/{case['id']}/documents", headers=users["inv1"],
                    json={"type_code": "victim_recognition", "values": {"harm": "выдуманный вред"}})
    assert r.status_code == 422


def test_special_procedure_warning(client, users, case):
    doc = client.post(f"/api/v1/cases/{case['id']}/documents", headers=users["inv1"],
                      json={"type_code": "information_request", "values": {"secrecy": "тайна связи"}}).json()
    items = doc["current_version"]["validation"]["items"]
    assert any(i["code"] == "special_procedure" and i["level"] == "error" for i in items)


def test_head_resolution_author_and_court_flags(client, users, case):
    doc = client.post(f"/api/v1/cases/{case['id']}/documents", headers=users["inv1"], json={"type_code": "custody_petition"}).json()
    assert doc["template"]["requires_court"] is True and doc["template"]["requires_head_consent"] is True
    assert "СОГЛАСЕН" in doc["current_version"]["content_html"]
    doc = client.post(f"/api/v1/cases/{case['id']}/documents", headers=users["inv1"], json={"type_code": "joinder"}).json()
    assert doc["template"]["author_role"] == "Руководитель следственного органа"
